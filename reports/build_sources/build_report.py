from pathlib import Path
from copy import deepcopy
from io import BytesIO
import csv,json,zipfile,hashlib,re,shutil
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[2]
BUILD=ROOT/'.build/report-authoring'
BUILD.mkdir(parents=True,exist_ok=True)
OUT=ROOT/'reports/final/BaoCao_Airline_RNN.docx'
OUT.parent.mkdir(exist_ok=True,parents=True)
REF=ROOT/'sources/report_template.docx'
GROUP=json.loads((ROOT/'airline_rnn/web/project-info.json').read_text())
D=Document(REF); body=D._element.body
original_sects=[deepcopy(s._sectPr) for s in D.sections]
cover=[deepcopy(p._p) for p in D.paragraphs[:22]]
for e in list(body):body.remove(e)
for e in cover:body.append(e)
body.append(original_sects[1])
# Retain all cover components, including original VML frame and embedded logo.
for old,new in [('TIỂU LUẬN CHUYÊN NGÀNH','TIỂU LUẬN HỌC PHẦN'),('PHÂN LOẠI ẢNH BẰNG GRAPH CONVOLUTIONAL NETWORK TRÊN ĐỒ THỊ SUPERPIXEL',GROUP['title'].upper()),('TÊN HỌC PHẦN: THỊ GIÁC MÁY TÍNH','TÊN HỌC PHẦN: DEEP LEARNING')]:
 for p in D.paragraphs:
  if p.text.strip()==old:
   runs=p.runs
   for r in runs:r.text=''
   runs[0].text=new
# Fill the four existing cover slots, keeping the original cover topology.
cover_lines=['NHÓM THỰC HIỆN']+[m['student_id']+' - '+m['name'].upper() for m in GROUP['members']]
for i,text in zip([11,12,13,14],cover_lines):
 z=D.paragraphs[i]
 for run in z.runs:run.text=''
 run=z.runs[0] if z.runs else z.add_run()
 run.text=text;run.font.name='Times New Roman';run.font.size=Pt(13);run.bold=i==11
 run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
 z.alignment=WD_ALIGN_PARAGRAPH.CENTER
 z.paragraph_format.line_spacing=1.1;z.paragraph_format.space_before=Pt(0);z.paragraph_format.space_after=Pt(2)
# A heading break explicitly starts first TOC page, avoiding source field caches.
headings=[];figs=[];tabs=[];content=[];toc_slots=[]
page_map=json.loads((Path(__file__).parent/'page-map.json').read_text()) if (Path(__file__).parent/'page-map.json').exists() else {}
def font(r,size=13,bold=None,italic=None):
 r.font.name='Times New Roman';r.font.size=Pt(size)
 r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
 if bold is not None:r.bold=bold
 if italic is not None:r.italic=italic
 return r
def p(text='',style='Normal',size=13,align=None,indent=None,before=0,after=5,keep=False):
 z=D.add_paragraph(style=style);font(z.add_run(text),size)
 f=z.paragraph_format;f.space_before=Pt(before);f.space_after=Pt(after);f.line_spacing=1.5;f.keep_with_next=keep;f.widow_control=True
 if align is not None:z.alignment=align
 if indent is not None:f.first_line_indent=Cm(indent)
 content.append({'type':'paragraph','text':text})
 return z
def bookmark(z,key):
 i=str(1000+len(headings)+len(figs)+len(tabs));st=OxmlElement('w:bookmarkStart');st.set(qn('w:id'),i);st.set(qn('w:name'),key);en=OxmlElement('w:bookmarkEnd');en.set(qn('w:id'),i);z._p.insert(1,st);z._p.append(en)
def heading(text,level=2,newpage=False):
 z=p(text,'Heading '+str(level),16 if level==1 else 14 if level==2 else 13,WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT,0,after=12 if level==1 else 5,keep=True)
 z.paragraph_format.page_break_before=newpage
 for r in z.runs:font(r,16 if level==1 else 14 if level==2 else 13,True)
 key='h'+str(len(headings)+1);bookmark(z,key);headings.append((text,level,key));return z
def plain_head(text,newpage=True):
 z=p(text,'Front Heading',16,WD_ALIGN_PARAGRAPH.CENTER,0,after=12,keep=True);font(z.runs[0],16,True);z.paragraph_format.page_break_before=newpage;return z
def toc_entry(text,key,size=11,bold=False,indent=0):
 z=p('',size=size,align=WD_ALIGN_PARAGRAPH.LEFT,indent=0,after=3)
 z.paragraph_format.line_spacing=1.1;z.paragraph_format.left_indent=Cm(indent)
 z.paragraph_format.tab_stops.add_tab_stop(Cm(15.5-indent),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),key)
 r=font(z.add_run(text),size,bold);z._p.remove(r._r);h.append(r._r);z._p.append(h)
 font(z.add_run('\t'+str(page_map.get(text,'—'))),size)
 return z
def table(title,headers,rows,widths=None,size=11.5):
 z=p(title,'Table Caption',12,WD_ALIGN_PARAGRAPH.CENTER,0,after=4,keep=True);font(z.runs[0],12,False,True)
 key='t'+str(len(tabs)+1);bookmark(z,key);tabs.append((title,key))
 t=D.add_table(rows=1, cols=len(headers));t.autofit=False
 from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
 t.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=widths or [15.5/len(headers)]*len(headers)
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'000000');borders.append(b)
 pr.append(borders)
 margins=OxmlElement('w:tblCellMar')
 for side,v in [('top',80),('bottom',80),('left',85),('right',85)]:
  b=OxmlElement('w:'+side);b.set(qn('w:w'),str(v));b.set(qn('w:type'),'dxa');margins.append(b)
 pr.append(margins)
 for vals in [headers]+[[str(c) for c in row] for row in rows]:
  cells=t.rows[0].cells if vals is headers else t.add_row().cells
  for j,(c,txt) in enumerate(zip(cells,vals)):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   z=c.paragraphs[0];z.paragraph_format.first_line_indent=Cm(0);z.paragraph_format.space_after=Pt(0);z.paragraph_format.line_spacing=1.05;z.paragraph_format.keep_with_next=False
   z.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
   font(z.add_run(str(txt)),size,vals is headers)
  trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 content.append({'type':'table','title':title,'headers':headers,'rows':rows})
 p('',size=6,indent=0,after=0).paragraph_format.line_spacing=1
 return t
def fig(stem,title,height=7.2):
 f=ROOT/'results/figures'/f'{stem}.png';im=Image.open(f);w,h=im.size
 width=min(15.5,height*w/h);z=p('',align=WD_ALIGN_PARAGRAPH.CENTER,indent=0,after=0,keep=True)
 z.paragraph_format.line_spacing=1;z.add_run().add_picture(str(f),width=Cm(width))
 z=p(title,'Figure Caption',12,WD_ALIGN_PARAGRAPH.CENTER,0,after=6);font(z.runs[0],12,False,True)
 key='f'+str(len(figs)+1);bookmark(z,key);figs.append((title,key));content.append({'type':'figure','path':str(f.relative_to(ROOT)),'title':title})
def mr(t):
 r=OxmlElement('m:r');pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);r.append(pr);tx=OxmlElement('m:t');tx.text=t;r.append(tx);return r
def sub(a,b):
 el=OxmlElement('m:sSub');x=OxmlElement('m:e');x.append(mr(a));y=OxmlElement('m:sub');y.append(mr(b));el.extend([x,y]);return el
def frac(a,b):
 el=OxmlElement('m:f');x=OxmlElement('m:num');y=OxmlElement('m:den')
 for container,items in [(x,a),(y,b)]:
  for item in (items if isinstance(items,list) else [items]):container.append(mr(item) if isinstance(item,str) else item)
 el.extend([x,y]);return el
def equation(items,num):
 z=p('',align=WD_ALIGN_PARAGRAPH.CENTER,indent=0,after=8)
 m=OxmlElement('m:oMath')
 for item in items:m.append(mr(item) if isinstance(item,str) else item)
 z._p.append(m);font(z.add_run('    ('+num+')'),12)
 content.append({'type':'equation','number':num})
def csvrows(name):return list(csv.DictReader((ROOT/'results/tables'/name).open()))
def num(v,n=4):return f'{float(v):.{n}f}'.replace('.',',')
def pc(v):return num(float(v)*100,2)+'%'
# Fixed front-matter slots will be filled once body headings/captions are known.
for name in ['MỤC LỤC','MỤC LỤC (tiếp theo)','DANH MỤC HÌNH','DANH MỤC BẢNG']:
 z=plain_head(name);toc_slots.append((name,z._p))
plain_head('DANH MỤC TỪ VIẾT TẮT')
table('Các thuật ngữ sử dụng trong báo cáo',['Viết tắt','Diễn giải'],[['RNN','Recurrent Neural Network'],['LSTM','Long Short-Term Memory; chỉ thuộc notebook nguồn'],['NLP','Natural Language Processing'],['OOV / PAD','Out of Vocabulary / Padding'],['CE','Cross-Entropy'],['Val','Validation, tập xác thực'],['SD','Độ lệch chuẩn giữa training seeds'],['API','Application Programming Interface'],['CPU / GPU','Central / Graphics Processing Unit']],[3.5,12])
# Preserve original first-section geometry/references/numbering in the break.
z=p('',size=1,indent=0,after=0);z.paragraph_format.line_spacing=1;z._p.get_or_add_pPr().append(original_sects[0])
heading('MỞ ĐẦU',1)
heading('Giới thiệu đề tài')
p('Phản hồi trên mạng xã hội chứa nhiều thông tin về trải nghiệm sử dụng dịch vụ. Với ngành hàng không, một tweet có thể là lời cảm ơn, câu hỏi về lịch bay hoặc phàn nàn về hành lý và chậm chuyến. Phân tích cảm xúc tự động giúp chuyển văn bản thành nhãn có thể thống kê, nhưng ngôn ngữ ngắn, viết tắt và thiếu ngữ cảnh khiến bài toán khó hơn việc dò các từ tích cực hay tiêu cực [1].')
p('Đề tài nghiên cứu mạng nơ-ron hồi quy thuần RNN cho phân loại cảm xúc tweet hàng không. Mô hình xử lý chuỗi token theo thứ tự, học embedding cùng bộ phân loại và trả về phân bố xác suất của ba lớp negative, neutral, positive. Giá trị của thực nghiệm nằm ở quy trình có thể chạy lại và bằng chứng phân tích theo lớp, bên cạnh kết quả Accuracy tổng thể.')
heading('Mục tiêu nghiên cứu')
p('Mục tiêu A là kiểm tra notebook Kaggle “Airline Sentiment Analysis (90% Accuracy) using RNN” [2] và tái hiện pipeline nguồn trong giới hạn mô hình RNN. Audit cho thấy notebook thực tế dùng LSTM và đã bỏ neutral. Vì vậy A1 được gọi là baseline chuyển thể từ nguồn, không phải tái lập nguyên trạng công bố 90%. A2 là nhánh binary sạch để quan sát độ ổn định khi áp dụng protocol tách dữ liệu và chọn checkpoint.')
p('Mục tiêu B là xây dựng benchmark ba lớp sạch, giữ các yếu tố cố định và khảo sát riêng độ dài chuỗi, số hidden units, class weighting. Mô hình cuối được chọn bằng validation Macro-F1 trước khi mở test. Báo cáo cung cấp bảng cấu hình, lịch sử huấn luyện, confusion matrix, phân tích lỗi và demo dùng đúng checkpoint đã chọn.')
heading('Đối tượng và phạm vi nghiên cứu')
p('Dữ liệu là Tweets.csv từ Twitter US Airline Sentiment [3], gồm 14.640 dòng gốc. Thực nghiệm chính giới hạn ở văn bản tiếng Anh và ba nhãn của dataset; các trường airline/tweet_id chỉ phục vụ audit, thống kê hoặc nhóm dữ liệu, không là feature đưa vào mô hình. Tất cả 22 lần huấn luyện của dự án sử dụng RNN. LSTM, GRU, BERT và Transformer không được huấn luyện trong ma trận này.')
p('Các kết quả trình bày là số liệu đã đo trên CPU máy cá nhân. Độ lệch chuẩn qua ba training seeds phản ánh biến thiên tối ưu trên cùng một split; không được hiểu thành khoảng tin cậy trên mọi tập dữ liệu. Phân tích định tính do trợ lý đọc mẫu và đề xuất giả thuyết; đánh giá độc lập của thành viên nhóm vẫn là bước cần bổ sung nếu muốn kết luận chắc hơn.')
heading('CHƯƠNG 1. TỔNG QUAN VÀ CƠ SỞ LÝ THUYẾT',1,True)
heading('1.1. Sentiment Analysis và dữ liệu tuần tự')
p('Sentiment Analysis trong phạm vi đề tài là phân loại mức câu/tweet: mỗi văn bản nhận đúng một nhãn trong ba lớp. Negative phản ánh trải nghiệm hoặc thái độ tiêu cực, positive phản ánh đánh giá tích cực, neutral chủ yếu truyền thông tin hay yêu cầu dịch vụ. Một câu vừa cảm ơn vừa nêu sự cố vẫn chỉ có một ground truth, nên ranh giới nhãn có thể không rõ khi thiếu chuỗi hội thoại [1].')
p('Thứ tự từ ảnh hưởng ý nghĩa. “Good service” khác “not good service”; “did not have to wait” lại có thể mô tả trải nghiệm tốt. RNN khai thác thứ tự bằng cách cập nhật trạng thái ẩn khi đọc từng token. Tweet có độ dài khác nhau cần được ánh xạ về tensor cùng kích thước để ghép batch, trong khi thông tin từ phần đệm cần được tách khỏi văn bản thật.')
heading('1.2. Tiền xử lý và biểu diễn chuỗi')
p('Tiền xử lý là ánh xạ xác định từ raw text sang clean text trước tokenization. Các phép lower-case, chuẩn hóa Unicode và thay URL/mention giúp giảm biến thể hình thức. Tuy nhiên việc bỏ dấu câu, emoji hoặc danh tính hãng có thể làm mất tín hiệu cảm xúc. Do đó quy tắc được ghi rõ, lưu song song raw/clean và áp dụng thống nhất khi train, đánh giá, inference; tác động hiệu năng chỉ được kết luận sau một đối chứng riêng.')
p('Tokenizer của pipeline chính tách từ theo khoảng trắng sau cleaning. Vocabulary là ánh xạ token sang chỉ số, được học chỉ từ train. Từ chưa gặp ở validation/test dùng OOV; PAD dùng để bổ sung chuỗi ngắn. Với T=40, padding và truncation ở cuối. Masking làm lớp hồi quy bỏ qua PAD. Cấu hình T=20 có thể cắt nội dung cuối, trong khi tăng T vượt độ dài thực chỉ tăng số vị trí đệm.')
heading('1.3. Word embedding')
p('Embedding là bảng vector có kích thước V × D. Mỗi token ID tra ra một vector D chiều, sau đó được cập nhật cùng mô hình bằng gradient. Đề tài dùng V=4.000 và D=128, không nạp embedding pretrained. ID 0 dành cho PAD, ID 1 dành cho OOV, các ID 2–4 dành cho marker. Keras hỗ trợ mask_zero để truyền mặt nạ của PAD tới lớp hồi quy [4].')
equation([sub('x','t'),' = E[',sub('w','t'),'] ∈ ℝ',mr('¹²⁸')],'1.1')
p('Ở đây w là token ID tại thời điểm t, E là ma trận embedding và x là vector đầu vào của bước hồi quy. Các ID chỉ là mã tra cứu, không mang quan hệ thứ tự số học. Mô hình học cách biểu diễn hữu ích cho nhiệm vụ trong phạm vi tập train, nên vector không tự bảo đảm hiểu được sarcasm hoặc các từ ngoài vocabulary.')
heading('1.4. Mạng nơ-ron hồi quy thuần RNN')
p('RNN duy trì một trạng thái ẩn H chiều. Tại mỗi bước, lớp kết hợp vector hiện tại với trạng thái trước qua các ma trận học được rồi áp dụng tanh. Cùng một bộ trọng số được dùng xuyên suốt chuỗi. Đề tài lấy trạng thái cuối của chuỗi đã masking để đại diện tweet, sau đó chuyển qua Dense và softmax [5].')
equation([sub('h','t'),' = tanh(',sub('W','x'),sub('x','t'),' + ',sub('W','h'),sub('h','t−1'),' + b)'],'1.2')
equation(['Parameters(RNN) = H(D + H + 1)'],'1.3')
p('Với D=128 và H=196, lớp hồi quy có 196 × (128 + 196 + 1) = 63.700 tham số. Khác với kiến trúc có cổng như LSTM/GRU, RNN chỉ có phép cập nhật trạng thái cơ bản. Việc tập trung một kiến trúc giúp các đối chứng phản ánh ảnh hưởng của biểu diễn, dung lượng và trọng số lớp trong cùng loại mô hình.')
heading('1.5. Huấn luyện, gradient và regularization')
p('Backpropagation Through Time lan truyền gradient qua các bước của chuỗi. Khi nhân lặp các đạo hàm, gradient có thể nhỏ dần hoặc tăng lớn. Pascanu, Mikolov và Bengio phân tích hai vấn đề này và đề xuất clipping chuẩn gradient cho trường hợp exploding gradient [6]. Trong pipeline chính, global clipnorm=1,0 giới hạn chuẩn của toàn bộ gradient; thao tác này không bảo đảm giải quyết vanishing gradient.')
p('Adam điều chỉnh cập nhật theo các ước lượng moment của gradient [7]. Learning rate trong dự án được khai báo tường minh là 0,001. Các lớp dropout giảm sự phụ thuộc vào một tập feature khi train [8]. SpatialDropout1D áp dụng trên embedding, dropout nội bộ của RNN tác động lên đầu vào/trạng thái hồi quy, còn dropout sau RNN và Dense tác động lên vector feature. Inference tắt dropout nên train fit Accuracy và validation Accuracy không cùng điều kiện đo.')
p('Cross-entropy phạt xác suất thấp của nhãn đúng. Với class weighting, mỗi mẫu train được nhân trọng số theo lớp để tăng ảnh hưởng lớp ít dữ liệu. Loss có trọng số của B5 không thể so trực tiếp về trị số với loss không trọng số của C0; quality được so bằng metric không trọng số trên validation chung.')
equation(['CE = − log ',sub('p','y'),' ;   ',sub('w','c'),' = ',frac([sub('N','train')],['K × ',sub('n','c,train')])],'1.4')
heading('1.6. Các metric đánh giá')
p('Confusion matrix có hàng là nhãn thật và cột là nhãn dự đoán. Với từng lớp c, TP là số dự đoán đúng lớp đó, FP là số mẫu lớp khác dự đoán thành c, FN là số mẫu c bị dự đoán sang lớp khác. Precision và Recall giúp phân biệt độ chính xác của dự đoán với mức thu hồi nhãn thật [9].')
equation([sub('Precision','c'),' = ',frac([sub('TP','c')],[sub('TP','c'),' + ',sub('FP','c')]),' ;   ',sub('Recall','c'),' = ',frac([sub('TP','c')],[sub('TP','c'),' + ',sub('FN','c')])],'1.5')
equation([sub('F1','c'),' = ',frac(['2 × ',sub('Precision','c'),' × ',sub('Recall','c')],[sub('Precision','c'),' + ',sub('Recall','c')])],'1.6')
p('Macro-F1 là trung bình F1 của tất cả lớp với trọng số bằng nhau. Weighted-F1 lấy trọng số support của từng lớp. Accuracy là tỷ lệ nhãn dự đoán đúng trên toàn bộ N mẫu. Phép tính F1 phải dùng prediction trên cả tập dữ liệu; trung bình F1 của các minibatch không tương đương F1 toàn tập [9]. Khi mẫu số bằng 0, dự án dùng zero_division=0 và vẫn tính đủ lớp.')
table('Bảng 1.1. Vai trò các metric trong thực nghiệm',['Metric','Vai trò'],[['Accuracy','Tỷ lệ tweet nhận đúng nhãn, cần đặt cạnh class prior'],['Precision / Recall / F1','Đánh giá từng sentiment và trade-off FP/FN'],['Macro-F1','Metric chính để checkpoint và chọn configuration'],['Weighted-F1','Chất lượng tổng hợp có xét support'],['Confusion matrix','Nhận diện cặp lớp thường nhầm'],['Classification report','Lưu đầy đủ per-class, macro/weighted average và support']],[4,11.5])
p('Dữ liệu gốc có negative chiếm khoảng 62,69%. Một mô hình thiên về lớp này có thể đạt Accuracy khá mà bỏ lỡ neutral/positive. Majority baseline ở benchmark chính đạt Accuracy 63,66% nhưng Macro-F1 chỉ 0,2593. Vì vậy báo cáo luôn phân biệt hiệu quả tổng thể với khả năng nhận diện đồng đều ba lớp.')
heading('CHƯƠNG 2. PHƯƠNG PHÁP VÀ GIẢI PHÁP',1,True)
heading('2.1. Pipeline tổng thể và nguyên tắc thực nghiệm')
p('Pipeline gồm kiểm tra nguồn, kiểm tra chất lượng dữ liệu, đóng băng split theo nhóm trùng, fit vocabulary trên train, biểu diễn chuỗi, huấn luyện từng configuration, chọn bằng validation và đánh giá test cuối. Tất cả run lưu cấu hình, seed, environment, checkpoint, history, thời gian và fingerprint của dữ liệu/mã nguồn. Các bước phân tích sau test không thay đổi lựa chọn mô hình.')
fig('F08_pipeline','Hình 2.1. Pipeline dữ liệu, huấn luyện và đánh giá với test holdout',5)
heading('2.2. Audit notebook Kaggle và định nghĩa baseline')
p('Notebook nguồn [2], phiên bản scriptVersionId=120745701, dùng text làm input và airline_sentiment làm nhãn. Sau khi loại 3.099 neutral và 5.000 negative đầu theo thứ tự nguồn, còn 6.541 tweet: 4.178 negative, 2.363 positive. Việc chọn mẫu theo thứ tự làm thay đổi phân bố hãng/lớp, không phải downsample ngẫu nhiên. Các cột khác dùng cho EDA, không đưa vào mạng.')
table('Bảng 2.1. Các thành phần xác định được từ notebook nguồn',['Thành phần','Giá trị/nhận xét đã audit'],[['Làm sạch','lstrip tập ký tự hãng; rstrip @; lower; regex [^a-zA-z0-9\\s]'],['Tokenizer / vocabulary','Keras Tokenizer, num_words=4.000; fit trên toàn 6.541 mẫu; không OOV'],['Chuỗi','Độ dài tự lấy max=31; pre-padding và pre-truncation'],['Label','negative=0, positive=1; sparse integer labels'],['Split','70/30 stratified, random_state=1; validation_split=0,2 từ phần train'],['Số mẫu train / val / test','3.662 / 916 / 1.963'],['Kiến trúc thực tế','Embedding(4.000,128), SpatialDropout(0,5), LSTM(196)'],['Các lớp sau hồi quy','Dropout(0,2), Dense(100,ReLU), Dropout(0,4), Dense(2,softmax)'],['Dropout nội bộ LSTM','input=0,3; recurrent=0,3; activation không khai báo tường minh'],['Train','Sparse CE; Adam; batch=32; 20 epoch; metric Accuracy'],['LR và runtime lịch sử','Không khai báo LR/version đầy đủ; chưa xác định giá trị default lịch sử'],['Checkpoint / Early Stopping','Không có lựa chọn checkpoint hoặc dừng sớm'],['Kết quả lưu trong notebook','Test Accuracy=90,32094%; test loss=0,6523001']],[4.2,11.3],11)
p('Output nguồn còn ghi train Accuracy epoch cuối 99,62%, validation Accuracy epoch cuối 90,61% và validation Accuracy tốt nhất 92,58% ở epoch 11. Đây là output lịch sử đã kiểm tra, không phải số nhóm đo lại. Test Accuracy được tính trên test binary đã lọc; tiêu đề có thể gây hiểu nhầm nếu coi đó là RNN hoặc ba lớp trên toàn dataset.')
p('Nguồn fit tokenizer trước split nên vocabulary biết phân bố token của test dù không dùng nhãn test trong fit. Audit nhánh A1 còn tìm thấy 68 tweet_id và 85 clean text trùng giữa các phần dữ liệu. Đây là rủi ro leakage; số overlap không xác định được Accuracy đã bị tăng bao nhiêu. Hàm lstrip dùng tập ký tự, không gỡ đúng tiền tố hãng như một chuỗi, vì vậy được giữ trong A1 để ghi nhận hành vi nguồn.')
p('A1 thay riêng lớp hồi quy bằng RNN(196), giữ source filtering/cleaning/tokenization/split và checkpoint epoch cuối. Adam LR=0,001 là quyết định triển khai tường minh, không khẳng định trùng default nguồn. A2 dùng binary dataset sạch và train-only tokenizer, vẫn giữ source cleaner, pre-padding, không masking/clipping; bổ sung checkpoint validation và Early Stopping. A1/A2 khác nhiều thành phần protocol nên không được dùng như một ablation một yếu tố để lượng hóa leakage.')
fig('F02_source_filtering','Hình 2.2. Việc lọc của nguồn biến bài toán ba lớp thành binary',6)
heading('2.3. Dataset và kiểm tra chất lượng')
table('Bảng 2.2. Đặc điểm dataset sử dụng',['Thuộc tính','Giá trị'],[['Nguồn','crowdflower/twitter-airline-sentiment; Tweets.csv'],['Dữ liệu gốc','14.640 dòng, 15 cột'],['Nhãn gốc','negative 9.178; neutral 3.099; positive 2.363'],['Input / output','Một tweet text / ba xác suất và nhãn argmax'],['Cột vào mô hình','text; label từ airline_sentiment'],['Cột audit/EDA','tweet_id, airline và raw/clean text'],['Loại khỏi benchmark chính','36 exact duplicates; 251 dòng thuộc nhóm xung đột nhãn'],['Dữ liệu sạch chính','14.353 dòng'],['Lưu provenance','Raw CSV, SHA-256, manifest, exclusion log và split indices']],[4.3,11.2])
p('Các bản sao hoàn toàn được loại trước. Các nhóm có cùng tweet_id, văn bản canonical hoặc clean text được nối bắc cầu; nhóm nhãn xung đột được cách ly, không tự suy đoán nhãn đúng. Nhóm còn lại đi trọn vào một split. Cách làm này kiểm soát duplicate leakage nhưng không loại được mọi quan hệ ngữ nghĩa gần nhau hay hội thoại liên tiếp. Log exclusions lưu row_id để truy vết 287 dòng loại.')
heading('2.4. Preprocessing chính P0')
p('P0 chuẩn hóa Unicode NFKC, giải mã HTML entities, chuẩn hóa apostrophe cong và chuyển lowercase. URL và mention được thay bằng urlmarker/usermarker. Các contraction phổ biến được mở rộng để giữ not, chẳng hạn can’t thành can not. Regex giữ chữ cái Latin, chữ số và khoảng trắng; dấu câu và emoji bị loại. Hashtag giữ nội dung từ khi ký tự # bị bỏ. Không bỏ stopwords, không stemming hoặc lemmatization. Văn bản rỗng sau cleaning được thay bằng emptymarker.')
p('P0 được định nghĩa trước chọn mô hình và áp dụng y hệt khi inference. Reserved IDs gồm PAD=0, OOV=1, usermarker=2, urlmarker=3, emptymarker=4. Train-only vocabulary có cap 4.000 tính cả reserved IDs; thứ tự token là xác định. Encoding nhãn chính cố định negative=0, neutral=1, positive=2. Tokenizer và label_map được lưu để demo không fit lại.')
fig('F07_preprocessing','Hình 2.3. Ví dụ raw text và clean text của P0',8.1)
p('Các thay đổi có hai mặt: mở rộng phủ định giữ thông tin not, nhưng gộp mention làm mất mục tiêu đánh giá trong câu so sánh hai hãng, bỏ ?? có thể mất sắc thái nghi vấn và bỏ emoji có thể mất thái độ. Các ví dụ sau test chỉ hỗ trợ chỉ ra thông tin đã mất, không chứng minh một quy tắc cleaning là nguyên nhân của chênh lệch metric. Nếu khảo sát P0 khác, cần ma trận mới và holdout chưa dùng.')
heading('2.5. Kiến trúc RNN và tham số')
fig('F09_architecture','Hình 2.4. Kiến trúc RNN ba lớp dùng trong C0 và B5',8.3)
arch=csvrows('final_architecture.csv')
table('Bảng 2.3. Kiến trúc RNN cuối; B là batch size',['Layer','Input','Output','Parameters'],[[r['type'],r['input_shape'].replace('None','B'),r['output_shape'].replace('None','B'),f"{int(r['parameters']):,}".replace(',','.') ] for r in arch],[4.2,3.8,4.2,3.3],11)
p('Embedding có 512.000 tham số, RNN có 63.700, Dense feature có 19.700 và Dense output có 303, tổng 595.703. Lớp RNN dùng tanh và return_sequences=False. Dense feature dùng ReLU, output dùng softmax. Spatial dropout=0,5; dropout input/recurrent=0,3; dropout sau RNN=0,2; dropout sau Dense=0,4. Mask PAD được bật trong các run ba lớp.')
heading('2.6. Training pipeline và lựa chọn mô hình')
table('Bảng 2.4. Cấu hình cố định của benchmark ba lớp',['Hyperparameter','Giá trị'],[['Training seeds','42, 2026, 3407; split seed=42'],['Vocabulary / embedding','4.000 / 128; fit vocabulary chỉ trên train'],['Chuỗi / padding / mask','T=40; post/post; mask_zero=True'],['Hidden / Dense','196 tanh / 100 ReLU'],['Optimizer / LR','Adam / 0,001'],['Loss / batch size','Sparse categorical cross-entropy / 32'],['Global gradient clipnorm','1,0'],['Epoch tối đa / patience','20 / 5; min_delta=0,001'],['Checkpoint','Validation Macro-F1 lớn nhất; nếu bằng thì val loss thấp hơn'],['Class weights C0 / B5','Không / N_train ÷ (3 × số mẫu lớp trên train)'],['Test gate','Chỉ mở sau frozen validation selection']],[5,10.5])
p('Python, NumPy và TensorFlow được đặt seed; TensorFlow op determinism được bật. Mỗi epoch lưu loss/Accuracy trong fit cùng train_eval Macro-F1 và validation Macro-F1 ở chế độ inference trên toàn tập. Early Stopping dùng mức cải thiện 0,001 và patience=5; checkpoint vẫn theo cực đại validation Macro-F1 chính xác, vì vậy epoch checkpoint và mốc tham chiếu Early Stopping có thể khác nhau.')
p('Configuration thắng được chọn theo mean validation Macro-F1 qua ba seeds. Với ứng viên cách top dưới 0,005, rule ưu tiên ít tham số, rồi median training wall time, rồi thứ tự matrix. B5-balanced thắng và seed 3407 là run có validation Macro-F1 trung vị, được chọn làm demo trước test. Model cuối sao chép checkpoint của run này, không refit train+validation để tránh tạo run mới chưa có validation.')
heading('CHƯƠNG 3. THỰC NGHIỆM VÀ ĐÁNH GIÁ',1,True)
heading('3.1. Môi trường và quy mô chạy')
p('Toàn bộ ma trận chạy trên máy cá nhân macOS arm64, CPU 10 physical/logical cores, RAM 16 GiB. TensorFlow không phát hiện GPU được dùng. Runtime là Python 3.12.14, TensorFlow 2.20.0, Keras 3.11.3, NumPy 2.2.6, pandas 2.3.3, scikit-learn 1.7.2. Dtype=float32; TensorFlow inter/intra threads cấu hình 1/2. Thông tin environment được lưu trong mỗi run.')
p('Tổng 22 run gồm A1 một seed, A2 ba seeds và sáu configurations ba lớp mỗi loại ba seeds. Tổng training wall time khoảng 19,91 phút, có tính bước tính metric cuối epoch. Đây là chi phí đo trên máy và môi trường trên; không phải thời gian GPU hay tổng thời gian chuẩn bị báo cáo. Kaggle/Colab là lựa chọn chuyển môi trường, chưa được chạy kiểm chứng cho gói hiện tại.')
heading('3.2. EDA, imbalance và split sạch')
fig('F01_raw_sentiment','Hình 3.1. Phân bố sentiment trên 14.640 tweet gốc',6)
fig('F03_airlines','Hình 3.2. Phân bố tweet theo hãng ở dữ liệu gốc',6)
p('Phân bố lớp lệch mạnh về negative. Phân bố theo airline cung cấp bối cảnh nguồn dữ liệu và giúp nhận diện ảnh hưởng của việc bỏ các negative đầu theo thứ tự. Airline không được sử dụng làm feature. Báo cáo không diễn giải chất lượng dịch vụ hiện tại của hãng từ mẫu tweet này.')
table('Bảng 3.1. Phân bố benchmark ba lớp sau kiểm tra dữ liệu',['Split','Negative','Neutral','Positive','Tổng'],[['Train',6404,2095,1544,10043],['Validation',1368,458,335,2161],['Test',1368,448,333,2149],['Tổng',9140,3001,2212,14353]],[3.5,3,3,3,3])
p('Tỷ lệ gần 70/15/15 được thực hiện có stratification theo nhóm và lớp, seed split=42. Group isolation dẫn đến số mẫu không đúng phép làm tròn chia từng dòng nhưng gần giữ class prior. Không có nhóm trùng chéo split sạch. Test gồm 2.149 mẫu cố định được dùng cho C0 và B5, nên hai kết quả này có thể đối chiếu trong cùng benchmark.')
fig('F05_train_lengths','Hình 3.3. Phân bố độ dài tweet P0 trên train',6)
p('Độ dài token trên train có median=20, P90=26, P95=28, P99=30, max=34. Giá trị T=40 bao phủ mọi tweet train của biểu diễn hiện tại; T=60 chủ yếu bổ sung PAD. Ngưỡng phân tích lỗi theo độ dài lấy từ quartile train (13,20,24), không tính từ test để lựa chọn hyperparameter.')
heading('3.3. Baseline A1 và A2')
tests=csvrows('test_summary.csv');v=csvrows('validation_summary.csv');vd={r['config_id']:r for r in v};td={r['config_id']:r for r in tests if r['config_id']!='majority'}
table('Bảng 3.2. Baseline binary, các tập đánh giá khác nhau',['Run','N test','Val Macro-F1','Test Accuracy','Test Macro-F1'],[[k,td[k]['support'],num(vd[k]['val_macro_f1_mean'])+((' ± '+num(vd[k]['val_macro_f1_std'])) if k=='A2' else ''),pc(td[k]['test_accuracy_mean']),num(td[k]['test_macro_f1_mean'])+((' ± '+num(td[k]['test_macro_f1_std'])) if k=='A2' else '')] for k in ['A1','A2']],[2,2,4,3.5,4],11)
p('A1 đạt Test Accuracy 83,80% và Macro-F1 0,8249 trên test binary 1.963 mẫu, dùng epoch cuối như nguồn. Số này thấp hơn output nguồn 90,32% nhưng kiến trúc đã đổi LSTM thành RNN và runtime/LR lịch sử chưa xác định; không thể coi chênh lệch là lỗi tái lập hay bằng chứng RNN luôn kém hơn. A1 còn giữ các rủi ro leakage để ghi nhận pipeline nguồn.')
p('A2 đạt mean Test Accuracy 74,08% ± 11,12 điểm phần trăm, Macro-F1 0,6531 ± 0,2183. Seed 42 dừng ở epoch 6 với checkpoint epoch 1, F1 positive chỉ khoảng 0,0593; các seed khác tốt hơn rõ rệt. Biến thiên lớn là kết quả cần báo cáo, không loại seed yếu để làm số đẹp. A2 thay cả dữ liệu, tokenizer và cách chọn epoch, vì vậy chưa tách được nguyên nhân bất ổn hay định lượng tác động cleaning/leakage.')
heading('3.4. Experiment Matrix và các yếu tố khảo sát')
table('Bảng 3.3. Ma trận thực nghiệm đã chạy',['ID','Benchmark / thay đổi','Số seeds','Mục đích'],[['A1','Binary nguồn, thay LSTM bằng RNN','1','Baseline chuyển thể'],['A2','Binary sạch, train-only vocab + checkpoint','3','Baseline sạch'],['C0','Ba lớp sạch, T40 H196, không class weights','3','Mốc đối chứng chính'],['B1-20','Chỉ T=20','3','Khảo sát truncation/chi phí'],['B1-60','Chỉ T=60','3','Kiểm tra thêm PAD'],['B3-64','Chỉ H=64','3','Giảm dung lượng RNN'],['B3-128','Chỉ H=128','3','Dung lượng trung gian'],['B5-balanced','Chỉ class_weight=balanced','3','Nhận diện lớp thiểu số']],[2.5,6.5,1.5,5],11)
p('Trong các run B, split, seeds, cleaner P0, vocabulary 4.000, embedding 128, dropout, LR, batch size, checkpoint rule và epoch budget giữ nguyên so với C0. Câu hỏi về sequence length là mức cắt/ngữ cảnh có thay đổi quality hoặc cost không; câu hỏi hidden units là dung lượng hồi quy có đem lại lợi ích không; câu hỏi class weighting là Macro-F1 và F1 neutral/positive có cải thiện khi trọng số lớp thay đổi không. Metric lựa chọn vẫn là validation Macro-F1, còn thời gian và tham số ghi như thước đo chi phí.')
rows=[]
for k in ['C0','B1-20','B1-60','B3-64','B3-128','B5-balanced']:
 r=vd[k];rows.append([k,pc(r['val_accuracy_mean']),num(r['val_macro_f1_mean'])+' ± '+num(r['val_macro_f1_std']),f"{int(r['parameters']):,}".replace(',','.'),num(r['training_seconds_median'],2)])
table('Bảng 3.4. So sánh validation qua ba seeds; thời gian median (giây)',['Config','Val Accuracy','Val Macro-F1 ± SD','Parameters','Time (s)'],rows,[2.6,3.1,4.5,2.8,2.5],11)
fig('F20_configuration_comparison','Hình 3.4. Validation Macro-F1 của sáu configurations ba lớp',6.5)
heading('3.5. Ảnh hưởng của sequence length')
p('T=20 có mean validation Macro-F1 0,6722, cao hơn T=40 (0,6623) trong mẫu này, đồng thời median training time giảm từ 76,49 xuống 39,61 giây. Tuy nhiên SD ở T=20 là 0,0159, lớn hơn mốc 0,0017. Cắt chuỗi có thể bỏ phần đánh giá cuối tweet, nên chất lượng trung bình cao hơn trong ba run chưa cho phép kết luận ngữ cảnh ngắn luôn có lợi.')
p('T=60 và T=40 có metric validation giống nhau ở từng seed. Điều này phù hợp với việc tweet sau P0 ngắn hơn 40 token và mask bỏ PAD: T=60 không thêm nội dung văn bản. Median thời gian tăng lên 108,86 giây, tức thêm khoảng 42% so với C0. Vì vậy đối chứng này chủ yếu chứng minh chi phí phần đệm, không phải khả năng học phụ thuộc xa của RNN.')
heading('3.6. Ảnh hưởng của hidden units và class weighting')
p('H=64 giảm tham số từ 595.703 xuống 531.155 và median wall time còn 30,32 giây; mean validation Macro-F1 0,6625 gần C0. H=128 đạt 0,6724 với 558.099 tham số và 50,47 giây. Embedding chiếm phần lớn tham số nên giảm H không làm tổng tham số giảm theo cùng tỷ lệ với thời gian. Chưa có bằng chứng H lớn nhất là tốt nhất trong miền đã khảo sát.')
p('B5-balanced có mean validation Macro-F1 0,6926 ± 0,0115, cao nhất ma trận ba lớp, trong khi mean validation Accuracy chỉ 75,21%. Class weights được tính từ 10.043 mẫu train; negative≈0,5227, neutral≈1,5979, positive≈2,1682. Các lớp ít dữ liệu đóng góp nhiều hơn vào train loss. B5 được chọn bằng validation, không dựa vào test.')
fig('F21_quality_cost','Hình 3.5. Quan hệ giữa validation Macro-F1 và chi phí huấn luyện',6.5)
heading('3.7. Kết quả test của mô hình cuối')
rows=[]
for k in ['C0','B5-balanced']:
 r=td[k];rows.append([k,pc(r['test_accuracy_mean'])+' ± '+num(float(r['test_accuracy_std'])*100,2)+' đpt',num(r['test_macro_f1_mean'])+' ± '+num(r['test_macro_f1_std']),num(r['test_weighted_f1_mean']),r['support']])
rows.append(['Majority','63,66%','0,2593','0,4952',2149])
table('Bảng 3.5. Test ba lớp; mean ± SD qua ba seeds trên cùng split',['Model','Accuracy','Macro-F1','Weighted-F1','N'],rows,[2.8,4.3,4.5,2.5,1.4],11)
p('B5-balanced đạt mean Test Accuracy 75,18% ± 1,34 điểm phần trăm; Macro-F1 0,6914 ± 0,0158, Weighted-F1 trung bình 0,7568. C0 đạt Accuracy 74,75% và Macro-F1 0,6515. Mức chênh Macro-F1 trung bình khoảng 0,0399 được quan sát trên cùng holdout. Ba seeds chưa đủ để khẳng định ý nghĩa thống kê rộng hay hiệu quả trên hãng/ngôn ngữ mới.')
p('Checkpoint seed 3407 dùng trong demo có Test Accuracy 75,38% và Macro-F1 0,6869. Những hình history, confusion matrix, ví dụ và bảng per-class dưới đây đều thuộc checkpoint này, không phải trung bình hoặc ensemble của ba seeds. Seed demo có validation Macro-F1 0,6935, checkpoint ở epoch 16 trong 20 epoch đã chạy.')
fig('F11_loss','Hình 3.6. Training Loss và Validation Loss của B5 seed 3407',6)
fig('F12_accuracy','Hình 3.7. Training Accuracy và Validation Accuracy của B5 seed 3407',6)
fig('F13_macro_f1','Hình 3.8. Train-eval và validation Macro-F1, B5 seed 3407',6)
p('History giúp nhận diện hội tụ và vị trí checkpoint tốt nhất. Train fit dùng dropout và class weights còn validation dùng inference và loss không class weights, nên không suy diễn overfitting chỉ từ việc hai đường loss nằm trên/dưới nhau. Train-eval Macro-F1 tính inference toàn train cho phép so với validation trong điều kiện đo gần nhau hơn. Không lấy metric epoch cuối thay metric checkpoint nếu checkpoint đã chọn khác epoch.')
heading('3.8. Confusion matrix và kết quả theo lớp')
fig('F14_confusion_count','Hình 3.9. Confusion matrix số lượng, B5 seed 3407; hàng là nhãn thật',7.4)
fig('F15_confusion_normalized','Hình 3.10. Confusion matrix chuẩn hóa theo hàng, B5 seed 3407',7.4)
per=csvrows('final_per_class.csv')
table('Bảng 3.6. Classification report theo lớp, B5 seed 3407',['Sentiment','Precision','Recall','F1-score','Support'],[[r['sentiment'],num(r['precision']),num(r['recall']),num(r['f1-score']),r['support']] for r in per],[3.5,3,3,3,3],12)
p('Negative có F1 cao nhất 0,8481 với 1.153/1.368 mẫu đúng. Positive nhầm thành neutral 79/333 mẫu (23,72%), thành negative 41/333. Neutral nhầm negative 157/448 (35,04%), positive 37/448. Neutral có F1 thấp nhất 0,5342. Các nhận xét này được rút trực tiếp từ confusion matrix của checkpoint seed 3407; không khẳng định negative luôn dễ nhất trên mọi split.')
fig('F16_per_class_f1','Hình 3.11. F1 theo sentiment của checkpoint demo',5.5)
p('So C0 với B5 ở mean ba seeds, F1 negative giảm từ 0,8596 xuống 0,8438; neutral tăng từ 0,4754 lên 0,5410; positive tăng từ 0,6196 lên 0,6895. Class weighting tạo trade-off giữa lớp đa số và hai lớp ít dữ liệu. Accuracy ít thay đổi trong khi Macro-F1 tăng cho thấy giá trị của việc đánh giá cân bằng theo lớp.')
heading('3.9. Phân tích lỗi theo độ dài, phủ định và OOV')
table('Bảng 3.7. Error slices theo độ dài, B5 seed 3407',['Số token','N','Accuracy','Macro-F1','Neg / Neu / Pos'],[['≤13',577,'0,6880','0,6909','220 / 213 / 144'],['14–20',607,'0,7496','0,6896','386 / 122 / 99'],['21–24',499,'0,7796','0,6004','378 / 59 / 62'],['>24',466,'0,8133','0,5248','384 / 54 / 28']],[3,1.5,3,3,5],11.5)
p('Accuracy tăng ở nhóm dài nhưng Macro-F1 giảm. Nhóm >24 token gồm 384 negative và chỉ 28 positive, trong khi nhóm ≤13 phân bố đều hơn. Vì vậy sự khác biệt có thể chịu ảnh hưởng class mix, không chứng minh tweet dài giúp RNN hiểu tốt hơn. Không có tweet test bị truncation ở T=40, nên lỗi checkpoint cuối không thể quy cho việc cắt cuối chuỗi này.')
p('Flag phủ định định nghĩa bằng not/no/never trong clean text. Nhóm có flag gồm 701 mẫu, Accuracy≈0,8103 và Macro-F1≈0,5200, với 591 negative; nhóm không flag có 1.448 mẫu, Accuracy≈0,7265 và Macro-F1≈0,7013. Khác biệt phân bố lớp làm so sánh này không phải đối chứng nhân quả về negation. Nhóm có OOV (1.219 mẫu) có Macro-F1≈0,6716 so với ≈0,7028 ở 930 mẫu không OOV; chưa kiểm soát độ dài/chủ đề/class mix.')
heading('3.10. Phân tích định tính các dự đoán')
p('85 case được chọn gồm tối đa 10 lỗi mỗi confusion pair, 5 case đúng mỗi lớp và 10 case có top-2 margin thấp. Trợ lý đã đọc raw/clean và lưu candidate explanation trong assistant_error_review.csv. Đây là mẫu có chủ đích để khám phá loại lỗi, không phải mẫu ngẫu nhiên dùng ước tính tỷ lệ sarcasm toàn dataset. Hai cột reviewer của nhóm trong manual_error_review.csv vẫn để trống; chưa có agreement của hai người.')
table('Bảng 3.8. Ví dụ lỗi có bằng chứng; ID là row_id của dữ liệu gốc',['ID / thật → dự đoán','Trích đoạn raw text','Nhận xét có giới hạn'],[['9332 / neg → pos','awesome... Doors close in 2 minutes ... WTH?','Ứng viên mỉa mai: lời khen đối lập sự cố; cần review người'],['1537 / neg → pos','good try but @SouthwestAir got her here safer and sooner','Hai hãng/mục tiêu; mention đều thành usermarker'],['14299 / neg → pos','thankfully ... Really not impressed ... :(','Mixed sentiment; not còn nguyên sau cleaning'],['8960 / neu → pos','@JetBlue really caring??','Cleaning bỏ ??; chưa có ablation để đo tác động'],['12077 / neu → neg','u r horrible ... hungupNOHELP','Nhãn neutral cần kiểm tra; không tự sửa nhãn'],['4647 / neu → pos','continues to prove to be the best airlines 💪','Ứng viên nhãn chưa phù hợp khi đọc riêng']],[3.2,6.5,5.8],11)
p('Một số tweet khó chứa lời cảm ơn xen phàn nàn, thành ngữ, so sánh đối tượng hoặc thông tin hội thoại bị thiếu. Tweet 9332 là ví dụ cụ thể cho khả năng sarcasm, nhưng chưa chứng minh mô hình hỏng mọi sarcasm. Tweet 14299 giữ not nhưng vẫn sai, còn tweet 10391 (“not ... truthful”) đúng: giữ phủ định là cần thiết về thông tin song không bảo đảm dự đoán đúng. Không sử dụng saliency/attribution trong lần chạy này, nên nhận xét về từ mà mô hình “bám vào” chỉ là giả thuyết.')
fig('F17_correct_examples','Hình 3.12. Các tweet dự đoán đúng của checkpoint cuối',8)
fig('F18_error_examples','Hình 3.13. Các tweet dự đoán sai của checkpoint cuối',8)
fig('F19_low_confidence','Hình 3.14. Các tweet có top-2 margin thấp',8)
p('Softmax confidence cao vẫn có thể sai: tweet 2856 được dự đoán positive với xác suất khoảng 92,66% dù ground truth negative. Ngược lại margin thấp không đồng nghĩa sai nhãn. Phân bố xác suất chưa được calibration; chưa báo ECE hay tỷ lệ từ chối. Vì vậy demo hiển thị xác suất do mô hình tính, không gọi đó là xác suất đúng đã được kiểm chứng [10].')
heading('3.11. So sánh với notebook gốc và hạn chế')
table('Bảng 3.9. Điều kiện so sánh nguồn và thực nghiệm chính',['Tiêu chí','Notebook nguồn','Dự án B5'],[['Loại mô hình','LSTM196','RNN196'],['Số lớp / dữ liệu','2 lớp; 6.541 mẫu lọc','3 lớp; 14.353 mẫu sạch'],['Fit tokenizer','Toàn corpus trước split','Chỉ train'],['Duplicate isolation','Không kiểm soát trong source pipeline','Nhóm trùng không chéo split'],['Chọn epoch','Epoch cuối (20)','Checkpoint validation Macro-F1'],['Kết quả test','90,32% lịch sử; N=1.963','75,18% ± 1,34 đpt; N=2.149'],['Macro-F1 / repeat seeds','Không thấy báo đủ trong nguồn','0,6914 ± 0,0158; 3 seeds']],[3.8,5.8,5.9],11)
p('Hai mức Accuracy không đo cùng nhiệm vụ hoặc protocol. Báo cáo không tuyên bố tái tạo chính xác 90%, không xếp hạng RNN/LSTM bằng chênh lệch này và không suy ra phần giảm do bỏ leakage. Mục tiêu đã thực hiện là tái hiện có ghi nhận khác biệt, sau đó tổ chức thực nghiệm RNN độc lập có kiểm soát.')
p('Hạn chế gồm một split chính cố định, chỉ ba training seeds, một dataset/miền tiếng Anh, số yếu tố khảo sát hữu hạn và neutral còn yếu. Duplicate grouping không bảo đảm loại mọi liên hệ ngữ nghĩa. Label confidence không được dùng làm trọng số hoặc feature. Rà soát định tính chưa có người adjudicate; softmax chưa hiệu chỉnh; cloud/GPU chưa kiểm chứng. Các kết quả đều được giữ, gồm run A2 thất thường.')
heading('CHƯƠNG 4. DEMO VÀ KHẢ NĂNG TÁI LẬP',1,True)
heading('4.1. Chức năng và checkpoint sử dụng')
p('Demo local cho phép nhập một câu/tweet tiếng Anh và nhận sentiment dự đoán cùng xác suất negative, neutral, positive. Bộ inference tải models/final/model.keras, tokenizer.json, label_map.json và representation/config manifest. Demo dùng B5-balanced seed 3407, checkpoint epoch 16. Hash model cuối khớp checkpoint đã đánh giá; không huấn luyện hoặc fit tokenizer khi người dùng nhập câu.')
fig('F23_demo_report','Hình 4.1. Ảnh chụp demo thực tế với một câu cảm ơn minh họa',10.8)
p('Ảnh trên minh họa một request thành công, không bổ sung mẫu vào test hoặc chứng minh Accuracy chung. Inference đi qua đúng P0, chuyển ID, post-padding/truncation T40, chạy softmax và chọn argmax. Các trường confidence/margin hoặc OOV nếu có chỉ mô tả request hiện tại. Câu ngoài miền hoặc bằng tiếng Việt có thể cho kết quả chưa đáng tin vì mô hình được đánh giá trên tweet tiếng Anh.')
heading('4.2. Giao diện, API và kiểm tra')
table('Bảng 4.1. Các endpoint và hành vi đã kiểm tra',['Endpoint','Chức năng / kiểm tra'],[['GET /','Logo trường, thông tin đề tài/nhóm, form và kết quả'],['GET /api/health','Thông tin model và trạng thái'],['GET /api/project','Thông tin trường, bài tiểu luận và ba thành viên'],['POST /api/analyze','Stream JSON từng khâu, trả vector/nhãn/xác suất'],['POST /api/predict','JSON có trường text; trả nhãn/xác suất'],['Input rỗng / JSON lỗi','HTTP 400; thông báo lỗi'],['Probabilities','Hữu hạn, trong [0,1], tổng gần 1'],['Parity inference','Khớp 12 prediction đã lưu trong kiểm tra artifact']],[5,10.5],12)
p('Website hiển thị logo/tên trường, khoa, đề tài, giảng viên và ba thành viên. Mỗi request hiển thị tám khâu: tiếp nhận văn bản, cleaning P0, token hóa, ánh xạ token thành ID, padding/masking, forward pass RNN, đọc softmax và chọn lớp bằng argmax. Có thể mở xem OOV, phần bị cắt và vector thật; vector không thể hiện mức đóng góp từ. Thời gian đo theo request, không phải benchmark latency. Demo chạy tại http://127.0.0.1:8765.')
heading('4.3. Cách chạy lại và cấu trúc sản phẩm')
p('README.md hướng dẫn tạo môi trường Python phù hợp, cài requirements đã pin, dùng dữ liệu CSV kèm theo hoặc lấy dataset bằng script có kiểm tra SHA-256. Notebook Airline_RNN_Complete.ipynb chạy các phase theo thứ tự; output hiện tại đã được thực thi top-to-bottom với các artifact của 22 run. Khi cấu hình/mã/dữ liệu thay đổi cần dùng output directory mới, không ghi đè provenance của thực nghiệm đã công bố.')
table('Bảng 4.2. Các thư mục và artifact để truy vết',['Thư mục / tệp','Nội dung'],[['data/raw, processed, splits','CSV gốc, exclusions, group/split manifests'],['configs/ và airline_rnn/','Ma trận cấu hình và pipeline RNN'],['models/runs và models/final','Checkpoint theo run; bộ inference cuối'],['results/metrics và predictions','Classification reports, JSON metrics, row-level predictions'],['results/figures và tables','23 hình PNG/PDF và bảng CSV'],['notebooks/','Notebook hoàn chỉnh và HTML preview'],['reports/final/','Báo cáo Word, slide và kịch bản thuyết trình'],['README.md / environment/','Hướng dẫn chạy, versions, kiểm tra protocol']],[6,9.5],11.5)
p('25 kiểm tra đã pass, gồm 11 kiểm tra protocol/holdout gate, 10 trường hợp trace/input và 4 kiểm tra tương thích sau đổi tên. Kiểm tra artifact xác nhận fingerprint của 22 run, không overlap trong split sạch, source tokenizer parity trên 6.541 tweet, tính lại metrics/confusion matrix của 10 test prediction files và inference parity. Các kiểm tra kỹ thuật tăng khả năng truy vết nhưng không thay thế đánh giá tổng quát hóa trên dữ liệu mới.')
heading('4.4. Kịch bản trình diễn và mở rộng')
p('Kịch bản khoảng một phút: giới thiệu checkpoint, nhập câu cảm ơn minh họa, đọc ba xác suất, sau đó dùng một case sai từ test để giải thích hạn chế. Người trình bày cần phân biệt nhãn được mô hình trả với nhãn tự đánh giá của câu demo. Không lựa chọn lại seed/config theo các câu thử trực tiếp, không dùng demo để điều chỉnh metric đã công bố.')
p('Các mở rộng cần rà nhãn bởi hai người và holdout mới. Dấu câu/emoji/target mention, calibration, clipping, dropout, LR, vocabulary và embedding dimension chưa có đối chứng trong ma trận hiện tại. So sánh LSTM/GRU/Transformer cần tách riêng với dataset/protocol/ngân sách công bằng; RNN vẫn là mô hình chính.')
heading('KẾT LUẬN',1,True)
p('Đề tài đã xây dựng một thực nghiệm Sentiment Analysis bằng RNN có quy trình kiểm tra nguồn, xử lý dữ liệu, chia nhóm trùng, huấn luyện lặp seeds, chọn mô hình trên validation và đánh giá cuối trên test. Audit xác định notebook Kaggle thực tế dùng LSTM và binary, nên A1 được trình bày là baseline chuyển thể thay vì khẳng định tái lập công bố 90%.')
p('Trong benchmark ba lớp, B5-balanced có validation Macro-F1 cao nhất trong ma trận đã thử và đạt mean Test Accuracy 75,18% ± 1,34 điểm phần trăm, Macro-F1 0,6914 ± 0,0158. Class weighting cải thiện F1 neutral/positive với trade-off ở negative. T60 không bổ sung ngữ cảnh so với T40 trong dataset đã làm sạch nhưng tăng chi phí; số hidden units nhỏ hơn có thể giảm thời gian mà không làm Macro-F1 giảm rõ trong miền khảo sát.')
p('Phân tích confusion matrix chỉ ra neutral là lớp khó nhất của checkpoint demo, positive thường nhầm neutral và negative có F1 cao nhất. Error slices chịu ảnh hưởng class mix, còn các ví dụ mixed sentiment, sarcasm và target ambiguity là giả thuyết định tính cần nhóm kiểm tra. Kết quả chưa chứng minh tổng quát hóa ngoài miền, ý nghĩa thống kê rộng hoặc softmax calibration.')
p('Sản phẩm gồm notebook, dataset/provenance, 22 run/checkpoints/history, bảng/hình/metrics, demo local và hướng dẫn chạy lại. Báo cáo và slide sử dụng cùng nguồn số liệu, không tạo Accuracy/F1 giả. Những mở rộng sau bàn giao cần protocol mới với holdout chưa dùng để duy trì tính đáng tin cậy của đánh giá.')
heading('TÀI LIỆU THAM KHẢO',1,True)
refs=[
('Bo Pang và Lillian Lee (2008), Opinion Mining and Sentiment Analysis, Foundations and Trends in Information Retrieval 2(1–2), 1–135.','https://www.cs.cornell.edu/home/llee/opinion-mining-sentiment-analysis-survey.html'),
('Chibuzor Okocha, Airline Sentiment Analysis (90% Accuracy) using RNN, Kaggle; version 1/scriptVersionId 120745701. Bản nguồn và output audit lưu trong sources/.','https://www.kaggle.com/code/chibuzorokocha/airline-sentiment-analysis-90-accuracy-using-rnn'),
('CrowdFlower, Twitter US Airline Sentiment, Kaggle dataset, Tweets.csv.','https://www.kaggle.com/datasets/crowdflower/twitter-airline-sentiment'),
('Keras, Embedding layer, API documentation. Version triển khai dự án: Keras 3.11.3.','https://keras.io/api/layers/core_layers/embedding/'),
('Keras, RNN layer, API documentation.','https://keras.io/api/layers/recurrent_layers/simple_rnn/'),
('Razvan Pascanu, Tomas Mikolov và Yoshua Bengio (2013), On the Difficulty of Training Recurrent Neural Networks, arXiv:1211.5063.','https://arxiv.org/abs/1211.5063'),
('Diederik P. Kingma và Jimmy Ba (2015), Adam: A Method for Stochastic Optimization, ICLR, arXiv:1412.6980.','https://arxiv.org/abs/1412.6980'),
('Nitish Srivastava và cộng sự (2014), Dropout: A Simple Way to Prevent Neural Networks from Overfitting, JMLR 15, 1929–1958.','https://jmlr.org/papers/v15/srivastava14a.html'),
('scikit-learn 1.7.2 documentation, f1_score; macro/weighted averaging và zero_division.','https://scikit-learn.org/1.7/modules/generated/sklearn.metrics.f1_score.html'),
('Chuan Guo, Geoff Pleiss, Yu Sun và Kilian Q. Weinberger (2017), On Calibration of Modern Neural Networks, ICML, PMLR 70, 1321–1330.','https://proceedings.mlr.press/v70/guo17a.html')]
for i,(t,url) in enumerate(refs,1):
 z=p(f'[{i}] {t}\n{url}',size=11.5,align=WD_ALIGN_PARAGRAPH.LEFT,indent=0,after=8);z.paragraph_format.line_spacing=1.05
p('Tài liệu trực tuyến đối chiếu ngày 06/10/2026. Số liệu thực nghiệm lấy từ artifact của dự án, không lấy từ các bài báo lý thuyết.',size=11.5,indent=0)
heading('PHỤ LỤC. TRUY VẾT VÀ THÔNG TIN NHÓM',1,True)
heading('A.1. Fingerprint và nguồn số liệu')
p('SHA-256 của Tweets.csv gốc: ea94b23f41892b290dec3330bb8cf9cb6b8bc669eaae5f3a84c40f7b0de8f15e. Model cuối: 94b36f13b045bebb406a11e753c40c32b83b6565bf3f0ddcde448144fc7cbba9. Các hash đầy đủ cho tokenizer/config/label map nằm trong models/final/manifest.json. Dùng hash để kiểm tra file chuyển máy không bị thay đổi.')
table('Bảng A.1. Ánh xạ số liệu báo cáo sang artifact',['Nội dung','Nguồn trong project'],[['Validation / test mean±SD','results/tables/validation_summary.csv; test_summary.csv'],['Per-class / CM demo','final_per_class.csv; B5-balanced_seed-3407_test.json'],['Split / chất lượng dữ liệu','dataset_splits.csv; exclusions.csv; data/splits/'],['History / epoch / time','models/runs/*/seed-*/history.csv và metadata'],['Ví dụ/qualitative hypotheses','manual_error_review.csv; assistant_error_review.csv'],['Mục đích 23 hình','results/figures/manifest.json'],['Protocol verification','results/metrics/artifact_verification.json']],[5,10.5],11)
heading('A.2. Thành viên và đóng góp thực tế')
p('Nhóm thực hiện gồm ba thành viên dưới đây. Phần đóng góp được điền theo công việc thực tế của từng thành viên.',size=12)
table('Bảng A.2. Danh sách thành viên nhóm',['Họ tên','MSSV','Đóng góp thực tế'],[[m['name'],m['student_id'],'[Nhóm điền]'] for m in GROUP['members']],[6,3,6.5],12)
# Fill source-styled native front matter with measured cached entries and internal hyperlinks.
entries=[(t,k,11 if lev==1 else 10.5,lev==1,0 if lev==1 else .45) for t,lev,k in headings if lev<=2 and not t.startswith('A.')]
# Two fixed TOC pages, split at balanced entry count.
for name,anchor in toc_slots:
 if name=='MỤC LỤC': chosen=entries[:len(entries)//2]
 elif name.startswith('MỤC LỤC ('):chosen=entries[len(entries)//2:]
 elif name=='DANH MỤC HÌNH':chosen=[(t,k,10.5,False,0) for t,k in figs]
 else:chosen=[(t,k,10.5,False,0) for t,k in tabs if t.startswith('Bảng')]
 cursor=anchor
 for vals in chosen:
  z=toc_entry(*vals);el=z._p;body.remove(el);cursor.addnext(el);cursor=el
body.append(original_sects[1])
# Save generated editable parts then merge with original package to keep opaque assets exact.
stage=BUILD/'report-generated.docx';D.save(stage)
with zipfile.ZipFile(REF) as src,zipfile.ZipFile(stage) as new,zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as out:
 edits={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','word/settings.xml'}
 names=set(src.namelist());newnames=set(new.namelist())
 for name in src.namelist():
  if name=='word/settings.xml':
   from lxml import etree
   settings=etree.fromstring(src.read(name));field=settings.find(qn('w:updateFields'))
   if field is None:field=OxmlElement('w:updateFields');settings.append(field)
   field.set(qn('w:val'),'true');out.writestr(name,etree.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True))
  else:out.writestr(name,new.read(name) if name in edits else src.read(name))
 for name in newnames-names:out.writestr(name,new.read(name))
with zipfile.ZipFile(REF) as src,zipfile.ZipFile(OUT) as result:
 preserved={n:hashlib.sha256(src.read(n)).hexdigest() for n in names-edits if result.read(n)==src.read(n)}
 assert len(preserved)==len(names-edits)
(BUILD/'report-index.json').write_text(json.dumps({'headings':headings,'figures':figs,'tables':tabs,'preserved_parts':len(preserved),'editable_parts':sorted(edits),'reference_sha256':hashlib.sha256(REF.read_bytes()).hexdigest()},ensure_ascii=False,indent=2))
(ROOT/'reports/final/report_content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2))
(ROOT/'reports/final/references.json').write_text(json.dumps([{'id':i,'title':t,'url':u,'access_date':'2026-10-06'} for i,(t,u) in enumerate(refs,1)],ensure_ascii=False,indent=2))
print(json.dumps({'path':str(OUT),'headings':len(headings),'figures':len(figs),'tables':len(tabs),'preserved_parts':len(preserved)},ensure_ascii=False))
